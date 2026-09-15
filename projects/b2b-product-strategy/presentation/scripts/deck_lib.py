"""Reusable design-system helpers matching the Nykaa-UX-Analysis.pptx style:
blush background, Cambria serif headlines, Calibri body, rose/green accent pair,
soft-shadow white cards, pill badges, dark-header tables, numbered circles."""

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

EMU_IN = 914400
SLIDE_W = 12192000
SLIDE_H = 6858000
MARGIN = 640080
CONTENT_W = SLIDE_W - 2 * MARGIN

BG = RGBColor(0xFA, 0xF5, 0xF3)
INK = RGBColor(0x1F, 0x16, 0x20)
ROSE = RGBColor(0xA8, 0x1E, 0x48)
GREEN = RGBColor(0x2F, 0x6F, 0x5E)
GRAY = RGBColor(0x6E, 0x5E, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROSE_LIGHT = RGBColor(0xF6, 0xDE, 0xE4)
GREEN_LIGHT = RGBColor(0xDC, 0xEC, 0xE6)
NEUTRAL_LIGHT = RGBColor(0xEC, 0xE2, 0xDF)
NEUTRAL = RGBColor(0x8A, 0x7B, 0x81)
LINE = RGBColor(0xDE, 0xD1, 0xCC)

F_HEAD = "Cambria"
F_BODY = "Calibri"

DECK_LABEL = "B2B Product Strategy"


def new_presentation():
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    return prs


def blank_layout(prs):
    return prs.slide_layouts[6]


def new_slide(prs):
    slide = prs.slides.add_slide(blank_layout(prs))
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(SLIDE_W), Emu(SLIDE_H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def _set_letter_spacing(run, pts):
    rPr = run._r.get_or_add_rPr()
    rPr.set('spc', str(int(pts * 100)))


def add_textbox(slide, text, x, y, w, h, size=11, color=GRAY, bold=False,
                 italic=False, font=F_BODY, align=PP_ALIGN.LEFT, spacing=0,
                 line_spacing=None, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
        if spacing:
            _set_letter_spacing(run, spacing)
    return tb


def add_rich_textbox(slide, x, y, w, h, runs_spec, align=PP_ALIGN.LEFT,
                      anchor=MSO_ANCHOR.TOP, line_spacing=None, wrap=True):
    """runs_spec: list of paragraphs, each a list of (text, dict(size,color,bold,italic,font))"""
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para_runs in enumerate(runs_spec):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for text, spec in para_runs:
            run = p.add_run()
            run.text = text
            run.font.size = Pt(spec.get('size', 11))
            run.font.bold = spec.get('bold', False)
            run.font.italic = spec.get('italic', False)
            run.font.name = spec.get('font', F_BODY)
            run.font.color.rgb = spec.get('color', GRAY)
    return tb


def add_eyebrow(slide, text, x=MARGIN, y=502920, w=CONTENT_W, color=ROSE, size=12, h=320040):
    return add_textbox(slide, text.upper(), x, y, w, h, size=size, color=color,
                        bold=True, font=F_BODY, spacing=2)


def add_headline(slide, text, x=MARGIN, y=822960, w=CONTENT_W, size=34, color=INK,
                  h=822960):
    return add_textbox(slide, text, x, y, w, h, size=size, color=color, bold=True,
                        font=F_HEAD, line_spacing=1.05)


def add_footer(slide, page_num, label=DECK_LABEL):
    add_textbox(slide, label, MARGIN, 6400800, 5486400, 320040, size=9.5, color=GRAY)
    add_textbox(slide, str(page_num), SLIDE_W - MARGIN - 914400, 6400800, 914400,
                320040, size=9.5, color=GRAY, align=PP_ALIGN.RIGHT)


def _round_rect(slide, x, y, w, h, adj=3183):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    shp.adjustments[0] = adj / 100000
    return shp


def add_card(slide, x, y, w, h, fill=WHITE, shadow=True, adj=3183, line_color=None):
    shp = _round_rect(slide, x, y, w, h, adj=adj)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color:
        shp.line.color.rgb = line_color
        shp.line.width = Emu(9525)
    else:
        shp.line.fill.background()
    if shadow:
        sp = shp._element.spPr
        effectLst = sp.makeelement(qn('a:effectLst'), {})
        shdw = sp.makeelement(qn('a:outerShdw'), {
            'sx': '100000', 'sy': '100000', 'kx': '0', 'ky': '0', 'algn': 'bl',
            'rotWithShape': '0', 'blurRad': '127000', 'dist': '38100', 'dir': '5400000'})
        clr = sp.makeelement(qn('a:srgbClr'), {'val': '1F1620'})
        alpha = sp.makeelement(qn('a:alpha'), {'val': '10000'})
        clr.append(alpha)
        shdw.append(clr)
        effectLst.append(shdw)
        sp.append(effectLst)
    shp.shadow.inherit = False
    return shp


def add_pill(slide, x, y, w, h, text, fill=GREEN, text_color=WHITE, size=11):
    shp = _round_rect(slide, x, y, w, h, adj=50000)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text.upper()
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.name = F_BODY
    run.font.color.rgb = text_color
    _set_letter_spacing(run, 1)
    return shp


def add_hline(slide, x, y, w, color=LINE, weight=9525):
    ln = slide.shapes.add_connector(1, Emu(x), Emu(y), Emu(x + w), Emu(y))
    ln.line.color.rgb = color
    ln.line.width = Emu(weight)
    return ln


def add_vline(slide, x, y, h, color=LINE, weight=9525):
    ln = slide.shapes.add_connector(1, Emu(x), Emu(y), Emu(x), Emu(y + h))
    ln.line.color.rgb = color
    ln.line.width = Emu(weight)
    return ln


def add_circle_badge(slide, x, y, d, text, fill=ROSE, text_color=WHITE, size=13):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x), Emu(y), Emu(d), Emu(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.name = F_BODY
    run.font.color.rgb = text_color
    return shp


def add_dot(slide, cx, cy, d, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx - d / 2)), Emu(int(cy - d / 2)), Emu(d), Emu(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _set_cell(cell, text, size=11, color=INK, bold=False, fill=None, align=PP_ALIGN.LEFT,
              font=F_BODY, anchor=MSO_ANCHOR.MIDDLE, margins=(91440, 91440, 45720, 45720)):
    cell.margin_left, cell.margin_right, cell.margin_top, cell.margin_bottom = (
        Emu(margins[0]), Emu(margins[1]), Emu(margins[2]), Emu(margins[3]))
    cell.vertical_anchor = anchor
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()
    tf = cell.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = color


def add_table(slide, x, y, w, col_widths, rows, row_heights=None, header=True,
              header_fill=INK, header_color=WHITE, zebra_fill=BG, body_size=10.5,
              header_size=10, first_col_bold=True, align_map=None):
    """rows: list of list-of-str. First row is header if header=True."""
    n_rows = len(rows)
    n_cols = len(col_widths)
    default_rh = 457200
    total_h = sum(row_heights) if row_heights else default_rh * n_rows
    gframe = slide.shapes.add_table(n_rows, n_cols, Emu(x), Emu(y), Emu(w), Emu(total_h))
    table = gframe.table
    # kill built-in banding style visuals by using plain style, then override
    tbl = table._tbl
    tblPr = tbl.find(qn('a:tblPr'))
    if tblPr is not None:
        tblPr.set('firstRow', '0')
        tblPr.set('bandRow', '0')
    for c, cw in enumerate(col_widths):
        table.columns[c].width = Emu(cw)
    for r in range(n_rows):
        table.rows[r].height = Emu(row_heights[r] if row_heights else default_rh)
        for c in range(n_cols):
            text = rows[r][c] if c < len(rows[r]) else ""
            is_header = header and r == 0
            align = PP_ALIGN.LEFT
            if align_map and c in align_map:
                align = align_map[c]
            if is_header:
                _set_cell(table.cell(r, c), text, size=header_size, color=header_color,
                          bold=True, fill=header_fill, align=align)
            else:
                fill = WHITE if (r % 2 == 1) else zebra_fill
                bold = first_col_bold and c == 0
                _set_cell(table.cell(r, c), text, size=body_size, color=INK if c == 0 else GRAY,
                          bold=bold, fill=fill, align=align)
    return gframe


def add_callout(slide, x, y, w, h, label, text, fill=NEUTRAL_LIGHT, label_color=ROSE,
                 text_color=INK, label_size=10.5, text_size=11.5):
    card = add_card(slide, x, y, w, h, fill=fill, shadow=False)
    pad = 228600
    inner_w = w - 2 * pad
    if label:
        add_rich_textbox(slide, x + pad, y, inner_w, h, [
            [(label.upper() + "   ", {'size': label_size, 'bold': True, 'color': label_color, 'font': F_BODY}),
             (text, {'size': text_size, 'bold': True, 'color': text_color, 'font': F_BODY})]
        ], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)
    else:
        add_textbox(slide, text, x + pad, y, inner_w, h, size=text_size, color=text_color,
                    bold=True, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)
    return card
