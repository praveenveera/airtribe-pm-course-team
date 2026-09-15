"""v2 deck design system: charcoal/amber, single grotesk sans throughout, flat
square-cornered cards with a left accent bar (no rounded pills, no serif) —
deliberately distinct from both v1's report template and the Nykaa-styled deck."""

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SLIDE_W = 12192000
SLIDE_H = 6858000
MARGIN = 640080
CONTENT_W = SLIDE_W - 2 * MARGIN

INK = RGBColor(0x1A, 0x1A, 0x1A)
AMBER = RGBColor(0xB5, 0x60, 0x12)
AMBER_DARK = RGBColor(0x8A, 0x48, 0x0B)
SLATE = RGBColor(0x5B, 0x64, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF2, 0xF1, 0xEF)
LIGHT_AMBER = RGBColor(0xFB, 0xEB, 0xD9)
LINE = RGBColor(0xE4, 0xE1, 0xDC)

FONT = "Arial"
DECK_LABEL = "B2B Product Strategy — v2"


def new_presentation():
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    return prs


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(SLIDE_W), Emu(SLIDE_H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def _spc(run, pts):
    rPr = run._r.get_or_add_rPr()
    rPr.set('spc', str(int(pts * 100)))


def add_textbox(slide, text, x, y, w, h, size=11, color=SLATE, bold=False, italic=False,
                 align=PP_ALIGN.LEFT, spacing=0, line_spacing=None, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = FONT
        r.font.color.rgb = color
        if spacing:
            _spc(r, spacing)
    return tb


def add_rich_textbox(slide, x, y, w, h, runs_spec, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs_spec):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for text, spec in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(spec.get('size', 11))
            r.font.bold = spec.get('bold', False)
            r.font.italic = spec.get('italic', False)
            r.font.name = FONT
            r.font.color.rgb = spec.get('color', SLATE)
    return tb


def add_eyebrow(slide, text, x=MARGIN, y=560000, w=CONTENT_W, color=AMBER, size=12):
    return add_textbox(slide, text.upper(), x, y, w, 320040, size=size, color=color, bold=True, spacing=2)


def add_headline(slide, text, x=MARGIN, y=900000, w=CONTENT_W, size=32, h=900000, color=INK):
    return add_textbox(slide, text, x, y, w, h, size=size, color=color, bold=True, line_spacing=1.05)


def add_footer(slide, page_num, label=DECK_LABEL):
    add_textbox(slide, label, MARGIN, 6420000, 5486400, 300000, size=9, color=SLATE)
    add_textbox(slide, str(page_num), SLIDE_W - MARGIN - 700000, 6420000, 700000, 300000,
                size=9, color=SLATE, align=PP_ALIGN.RIGHT)
    add_hline(slide, MARGIN, 6360000, CONTENT_W, color=LINE, weight=6350)


def add_hline(slide, x, y, w, color=LINE, weight=9525):
    ln = slide.shapes.add_connector(1, Emu(x), Emu(y), Emu(x + w), Emu(y))
    ln.line.color.rgb = color
    ln.line.width = Emu(weight)
    return ln


def add_vline(slide, x, y, h, color=LINE, weight=9525):
    ln = slide.shapes.add_connector(1, Emu(x), Emu(y), Emu(x), Emu(y + h))
    ln.line.color.rgb = color
    ln.line.width = Emu(weight)
    return ln


def add_rect(slide, x, y, w, h, fill=WHITE, line_color=None, line_w=9525):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color:
        shp.line.color.rgb = line_color
        shp.line.width = Emu(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_card(slide, x, y, w, h, fill=LIGHT_GRAY, accent=AMBER, accent_w=45720):
    """Flat, square-cornered card with a left accent bar — the v2 signature
    shape, replacing v1's rounded-and-shadowed card."""
    add_rect(slide, x, y, w, h, fill=fill)
    add_rect(slide, x, y, accent_w, h, fill=accent)
    return x + accent_w + 137160  # returns usable inner-left x


def add_tag(slide, x, y, w, h, text, fill=AMBER, text_color=WHITE, size=10):
    shp = add_rect(slide, x, y, w, h, fill=fill)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text.upper()
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = text_color
    _spc(r, 1)
    return shp


def add_number_badge(slide, x, y, d, text, fill=INK, text_color=WHITE, size=15):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(d), Emu(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = text_color
    return shp


def _cell(cell, text, size=10, color=INK, bold=False, fill=None, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.MIDDLE, margins=(91440, 91440, 45720, 45720)):
    cell.margin_left, cell.margin_right, cell.margin_top, cell.margin_bottom = (
        Emu(margins[0]), Emu(margins[1]), Emu(margins[2]), Emu(margins[3]))
    cell.vertical_anchor = anchor
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()
    tf = cell.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = color


def add_table(slide, x, y, w, col_widths, rows, row_heights=None, header=True,
              header_fill=INK, header_color=WHITE, body_size=9.5, header_size=9):
    n_rows = len(rows)
    n_cols = len(col_widths)
    default_rh = 457200
    total_h = sum(row_heights) if row_heights else default_rh * n_rows
    gframe = slide.shapes.add_table(n_rows, n_cols, Emu(x), Emu(y), Emu(w), Emu(total_h))
    table = gframe.table
    tblPr = table._tbl.find(qn('a:tblPr'))
    if tblPr is not None:
        tblPr.set('firstRow', '0')
        tblPr.set('bandRow', '0')
    for c, cw in enumerate(col_widths):
        table.columns[c].width = Emu(cw)
    for r in range(n_rows):
        table.rows[r].height = Emu(row_heights[r] if row_heights else default_rh)
        for c in range(n_cols):
            text = rows[r][c] if c < len(rows[r]) else ""
            is_header = header and r == 0
            if is_header:
                _cell(table.cell(r, c), text, size=header_size, color=header_color, bold=True, fill=header_fill)
            else:
                fill = WHITE if (r % 2 == 1) else LIGHT_GRAY
                _cell(table.cell(r, c), text, size=body_size, color=INK if c == 0 else SLATE,
                      bold=(c == 0), fill=fill)
    return gframe
