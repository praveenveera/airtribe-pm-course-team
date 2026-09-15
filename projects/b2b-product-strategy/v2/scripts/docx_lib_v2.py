"""v2 report design system: charcoal/amber, grotesk sans, table-forward report.
Distinct from v1's report template and from the Nykaa-styled story deck."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.section import WD_SECTION

FONT = "Arial"
INK = RGBColor(0x1A, 0x1A, 0x1A)
AMBER = RGBColor(0xB5, 0x60, 0x12)
AMBER_DARK = RGBColor(0x8A, 0x48, 0x0B)
SLATE = RGBColor(0x5B, 0x64, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_AMBER = "FBEBD9"
LIGHT_GRAY = "F2F1EF"
MID_GRAY = "E4E1DC"

PAGE_NUM = [0]


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.font.color.rgb = INK
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT)
    return doc


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge in kwargs:
            el = OxmlElement(f'w:{edge}')
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), str(kwargs[edge].get('sz', 4)))
            el.set(qn('w:color'), kwargs[edge].get('color', 'DDDDDD'))
            tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for side, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tcPr.append(mar)


def add_eyebrow(doc, text, color=AMBER):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = FONT
    return p


def add_h1(doc, text, size=22):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = INK
    r.font.name = FONT
    return p


def add_h2(doc, text, size=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = AMBER_DARK
    r.font.name = FONT
    return p


def add_body(doc, text, size=10.5, color=INK, italic=False, bold=False, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.italic = italic
    r.font.bold = bold
    r.font.name = FONT
    return p


def add_bullets(doc, items, size=10.5):
    for item in items:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run("—  ")
        r.font.color.rgb = AMBER
        r.font.bold = True
        r.font.size = Pt(size)
        r.font.name = FONT
        r2 = p.add_run(item)
        r2.font.size = Pt(size)
        r2.font.color.rgb = INK
        r2.font.name = FONT


def add_callout(doc, label, text, bg=LIGHT_AMBER):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, bg)
    cell_margins(cell, top=160, bottom=160, left=220, right=220)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    if label:
        r = p.add_run(label.upper() + "   ")
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = AMBER_DARK
        r.font.name = FONT
    r2 = p.add_run(text)
    r2.font.bold = True
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = INK
    r2.font.name = FONT
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers, rows, col_widths=None, body_size=9.5, header_size=9):
    n_cols = len(headers)
    table = doc.add_table(rows=len(rows) + 1, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    # header
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        shade_cell(cell, "1A1A1A")
        cell_margins(cell)
        p = cell.paragraphs[0]
        r = p.add_run(h.upper())
        r.font.bold = True
        r.font.size = Pt(header_size)
        r.font.color.rgb = WHITE
        r.font.name = FONT
    for ri, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(ri + 1, c)
            shade_cell(cell, LIGHT_GRAY if ri % 2 == 1 else "FFFFFF")
            cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(val))
            r.font.size = Pt(body_size)
            r.font.color.rgb = INK if c == 0 else SLATE
            r.font.bold = (c == 0)
            r.font.name = FONT
    if col_widths:
        for c, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[c].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_rule(doc, color="B56012"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '18')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_footer_text(section, left_text):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    r = p.add_run(left_text)
    r.font.size = Pt(8.5)
    r.font.color.rgb = SLATE
    r.font.name = FONT
    # page number field on the right via tab stop
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.3))
    tab_run = p.add_run("\t")
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = "PAGE"
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run_el = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '17')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '5B6470')
    rpr.append(sz)
    rpr.append(color)
    run_el.append(rpr)
    run_el.append(fld_begin)
    run_el.append(instr)
    run_el.append(fld_end)
    p._p.append(run_el)


def add_page_break(doc):
    doc.add_page_break()
