from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path('/Users/praveenveera/Documents/work-offline/product')
DOCX = OUT / 'Airtribe_B2B_Product_Strategy_Assessment_Pack.docx'
NAVY = '17365D'
LIGHT = 'F3F6F9'
GRAY = 'D9D9D9'

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)

def borders(table):
    tblpr = table._tbl.tblPr
    border = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '6')
        el.set(qn('w:color'), GRAY)
        border.append(el)
    tblpr.append(border)

def cell_text(cell, value, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(value))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    borders(t)
    for i, value in enumerate(headers):
        cell_text(t.rows[0].cells[i], value, True, 'FFFFFF')
        shade(t.rows[0].cells[i], NAVY)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value)
            if ridx % 2:
                shade(cells[i], LIGHT)
    for row in t.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def p(doc, text):
    para = doc.add_paragraph(text)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.08

def h(doc, text, level=1):
    para = doc.add_heading(text, level=level)
    para.paragraph_format.keep_with_next = True

def bullet(doc, text):
    para = doc.add_paragraph(text, style='List Bullet')
    para.paragraph_format.space_after = Pt(2)

def number(doc, text):
    para = doc.add_paragraph(text, style='List Number')
    para.paragraph_format.space_after = Pt(2)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12)]:
    styles[name].font.name = 'Aptos'
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph('B2B Product Strategy Assessment', style='Title')
subtitle = doc.add_paragraph('AI Customer Support to Customer Operations Strategy')
subtitle.runs[0].font.size = Pt(16)
subtitle.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
p(doc, 'Purpose: a complete evidence-controlled strategy pack for the Airtribe assessment. It covers the decision, research plan, customer discovery, market evidence, strategic options, roadmap, metrics, business model, risks, and the final under-five-minute Loom narrative.')
p(doc, 'Working conclusion: the company should not lead with another generic chatbot. It should test whether a vendor-neutral, governed AI support-operations layer can help support-led B2B companies automate selected resolutions across their existing helpdesk and CRM systems.')
add_table(doc, ['Document status', 'Evidence status', 'Next decision'], [[
    'Working strategy and research pack',
    'Market evidence reviewed; customer interviews still required',
    'Validate the initial segment, urgent workflow, and willingness to pay'
]], [1.6, 2.5, 2.3])
p(doc, 'Important limitation: no user interviews have been conducted in this pack. Interview findings, quotes, sample sizes, market size, and performance targets must not be presented as facts until collected and verified.')
doc.add_page_break()

h(doc, '1 Assignment interpretation')
p(doc, 'The assignment asks the Product Manager to define a clear two-to-three-year strategy for an AI-powered customer-support company whose adoption is slower than expected. The challenge is to move from a broad chatbot proposition to a durable customer-operations business without jumping directly into a feature list.')
add_table(doc, ['Question', 'Required answer'], [
    ['What problem?', 'Support teams struggle to resolve requests efficiently when knowledge, customer context, and actions are spread across systems.'],
    ['Who?', 'Initial hypothesis: support-led B2B SaaS or technology-enabled services companies with growing demand and lean support operations.'],
    ['Why now?', 'AI agents are moving from answer generation toward workflow action, but trust, security, governance, and outcome proof remain buying constraints.'],
    ['Why this product?', 'Potentially because it can operate across existing systems instead of requiring a full helpdesk or CRM replacement.'],
    ['How does it scale?', 'Land with one measurable support workflow, expand to more workflows and teams, then broaden into customer operations only after repeatable value is proven.']
], [1.6, 4.8])

h(doc, '2 Decision memo')
add_table(doc, ['Decision item', 'Recommendation'], [
    ['Initial segment', 'Support-led B2B SaaS and technology-enabled services companies; narrow further after interviews.'],
    ['Primary user', 'Support agents and support-operations leads.'],
    ['Economic buyer', 'Head or VP of Support, Customer Operations, or COO.'],
    ['Initial job', 'Resolve repetitive, high-volume requests faster using trusted customer and product context.'],
    ['Positioning hypothesis', 'A governed AI support-operations layer that automates selected resolutions across existing systems.'],
    ['Initial product scope', 'One workflow, one segment, and one or two integrations; grounded assistance, triage, low-risk actions, human handoff, and auditability.'],
    ['Primary proof', 'Verified AI-assisted resolutions without increased reopens, incorrect actions, or customer harm.']
], [1.6, 4.8])
p(doc, 'Trade-off: this is narrower and slower to market than launching a broad chatbot, but it creates a better path to measurable value, trust, expansion, and defensibility. If interviews show that customers prefer native AI from their existing vendor, reconsider the separate-platform strategy.')

h(doc, '3 Evidence and confidence')
p(doc, 'The current evidence supports the problem direction, but not yet the exact beachhead or willingness to pay. This table prevents strategic language from overstating what is known.')
add_table(doc, ['Statement', 'Type', 'Confidence', 'Unknown'], [
    ['Customer context is fragmented across support, CRM, product, billing, and knowledge systems.', 'Market/product evidence', 'Medium-high', 'Which segment has the highest economic impact?'],
    ['Security, governance, and human control affect AI adoption.', 'Market evidence', 'High', 'Which controls are purchase gates?'],
    ['Existing vendors offer AI agents, actions, governance, analytics, and handoffs.', 'Competitor evidence', 'High', 'Where is the unresolved gap?'],
    ['Support-led B2B SaaS with lean teams is the best starting segment.', 'Strategic hypothesis', 'Low-medium', 'Is the pain urgent and budgeted?'],
    ['Customers will pay for a vendor-neutral control layer.', 'Business hypothesis', 'Low', 'Separate layer or incumbent add-on?'],
    ['Verified resolution can demonstrate value.', 'Measurement hypothesis', 'Medium', 'What definition aligns with finance?']
], [2.2, 1.4, 0.9, 1.8])

h(doc, '4 Market and competitor validation')
p(doc, 'Major platforms already describe AI as an agent that can use knowledge, take actions, escalate to humans, and improve through operational feedback. This validates the need for a stronger strategy, but it also makes generic positioning weak.')
add_table(doc, ['Competitor', 'Observed capability', 'Implication'], [
    ['Intercom Fin', 'Outcome pricing, procedures, support and sales roles, existing-helpdesk use, operational AI.', 'AI agent plus outcomes is already established.'],
    ['Zendesk AI', 'Actions across systems, governance, quality controls, human handoffs, learning loops.', 'Trust and resolution alone are not sufficient.'],
    ['Freshdesk Freddy AI', 'AI agents, copilot, insights, agentic workflows, routing, permissions, audit logs.', 'Simplicity and accessibility may be strong incumbent advantages.'],
    ['Salesforce Agentforce', 'CRM-native agents, action pricing, conversations, workflow automation.', 'Salesforce-centric enterprises may not need a separate layer.']
], [1.2, 3.1, 2.3])
p(doc, 'Strategic implication: solve a specific cross-system problem that incumbents do not solve well, or focus on a narrow operational niche such as AI quality, governance, replay, and outcome measurement. Trusted AI resolution layer is a useful direction, not yet a proven moat.')

doc.add_page_break()
h(doc, '5 Customer discovery plan')
p(doc, 'The purpose of interviews is not to collect positive opinions about the proposed platform. It is to understand a recent support workflow, its failure points, the workaround, the business consequence, and the buying process.')
add_table(doc, ['Participant group', 'Target diversity', 'What to learn'], [
    ['Frontline support agents', 'Different industries and support maturity', 'Where agents search, switch tools, wait, correct AI, and escalate.'],
    ['Support or CS operations', 'Small and growing teams', 'Workflow ownership, reporting, routing, knowledge, and automation gaps.'],
    ['Support leader or COO', 'Different company sizes', 'Budget urgency, business outcomes, headcount pressure, and purchase triggers.'],
    ['IT, security, or procurement', 'Companies with formal controls', 'Permissions, data handling, auditability, vendor approval, and objections.']
], [1.6, 2.0, 3.0])
h(doc, 'Interview questions', 2)
for question in [
    'Tell me about the last difficult customer-support issue your team handled.',
    'Which systems and people were involved from intake to resolution?',
    'Where did the agent lose time or need to repeat information?',
    'What happens when volume increases unexpectedly?',
    'Which support metric is most painful or least trustworthy today?',
    'Which automation or AI tools are already in use? What still requires correction?',
    'Which actions would you never allow AI to take without approval?',
    'Tell me about the last support or AI tool purchase. What triggered it and who approved it?',
    'What evidence would be required to continue a pilot after 30 to 60 days?',
    'What would make you stop using an AI support tool?'
]: number(doc, question)
p(doc, 'Interview evidence rule: record the participant segment, role, company size band, industry, support maturity, workflow, pain, impact, workaround, and confidence. Anonymize notes. Do not invent or polish quotes.')

h(doc, '6 Segmentation and initial ICP hypothesis')
add_table(doc, ['Maturity', 'Operating reality', 'Likely need', 'Priority'], [
    ['Reactive/manual', 'Shared inboxes, spreadsheets, informal escalation.', 'Basic process, visibility, and repeatability.', 'Lower priority unless pain is severe.'],
    ['Standardized', 'Helpdesk, macros, basic SLA reporting, static knowledge base.', 'Cross-system context and workflow consistency.', 'Strong candidate.'],
    ['AI-assisted', 'Copilot, chatbot, routing, partial integrations.', 'Reliable knowledge, measurable ROI, safe actions.', 'Strong candidate.'],
    ['Integrated/agentic', 'AI connected to business systems with governance.', 'Continuous optimization and policy controls.', 'Later expansion or specialist niche.']
], [1.1, 2.0, 2.1, 1.2])
p(doc, 'ICP hypothesis: B2B SaaS or technology-enabled services companies with growing support volume, existing helpdesk and CRM investment, five to fifty support or customer-success staff, and a small operations function. These ranges are hypotheses for sampling, not market claims.')

doc.add_page_break()
h(doc, '7 Strategic options')
add_table(doc, ['Option', 'Value', 'Risk', 'Decision'], [
    ['A. Generic customer chatbot', 'Fast to explain and launch.', 'Low differentiation; incumbents already have it.', 'Do not prioritise.'],
    ['B. Embedded AI agent for one helpdesk', 'Fast time to value and lower integration scope.', 'Platform dependency and limited expansion.', 'Good pilot alternative.'],
    ['C. Vendor-neutral control layer', 'Addresses heterogeneous stacks and creates expansion potential.', 'Integration cost and incumbent competition.', 'Recommended only if validated.'],
    ['D. AI governance and quality layer', 'Clear trust and operational problem.', 'May be perceived as a feature.', 'Fallback or wedge if governance dominates.']
], [1.8, 1.8, 2.0, 0.8])
p(doc, 'Recommendation: validate Option C against Option B. If customers cannot articulate cross-system pain or will not buy a separate layer, choose the narrower embedded workflow option. This keeps learning fast and the decision reversible.')

h(doc, '8 Product strategy and roadmap')
add_table(doc, ['Stage', 'Strategic objective', 'Smallest product slice', 'Exit evidence'], [
    ['0–6 months: validate', 'Prove the wedge.', 'One workflow, one segment, one or two integrations; grounded assistance, triage, handoff, audit.', 'Repeated problem, willing pilots, no quality decline.'],
    ['6–12 months: pilot and harden', 'Expand trusted automation.', 'Low-risk actions, approval thresholds, permissions, analytics, fallback.', 'Measurable improvement and repeatable onboarding.'],
    ['Year 2: scale', 'Become the operational layer.', 'Multi-workflow orchestration, knowledge feedback, policy controls, broader integrations.', 'Usage expands across teams and workflows.'],
    ['Year 3: expand', 'Broaden customer operations.', 'Additional channels or adjacent workflows only where core value is proven.', 'Retention and expansion justify scope.']
], [1.4, 1.6, 2.5, 1.0])
h(doc, 'Non-goals', 2)
for item in [
    'Do not replace Zendesk, Salesforce, Freshdesk, or the customer CRM at launch.',
    'Do not build a proprietary foundation model.',
    'Do not support every channel, industry, and workflow before the wedge is proven.',
    'Do not allow high-risk actions such as refunds or account changes without explicit controls.',
    'Do not claim the platform strategy is validated before interviews and pilot evidence.'
]: bullet(doc, item)

h(doc, '9 AI product and enterprise controls')
add_table(doc, ['Control area', 'Minimum requirement'], [
    ['Identity and access', 'Tenant isolation, role-based permissions, least-privilege connector access, and customer-admin controls.'],
    ['Groundedness', 'Approved knowledge sources, source visibility, freshness monitoring, and refusal or escalation when context is missing or conflicting.'],
    ['Actions', 'Allowlisted tools, parameter validation, approval thresholds, idempotency, and prevention of irreversible actions by default.'],
    ['Safety', 'Prompt-injection defence, sensitive-data handling, policy enforcement, and handling of unauthorized requests.'],
    ['Human control', 'Edit, approve, reject, escalate, pause, and disable automation.'],
    ['Operations', 'Run identity, audit logs, latency, cost, provider failure, retry, fallback, and support telemetry.'],
    ['Rollout', 'Shadow mode or simulation, limited tenant pilot, stop rule, rollback trigger, and post-pilot review.']
], [1.5, 4.9])

h(doc, '10 Metrics and evaluation')
p(doc, 'The metric system should connect customer value, adoption, AI quality, safety, operations, and economics. Do not set numeric targets until baseline data exists.')
add_table(doc, ['Layer', 'Metrics'], [
    ['North-star candidate', 'Verified AI-assisted resolutions per active account.'],
    ['Customer value', 'Resolution time, backlog, SLA attainment, cost per resolved request, CSAT.'],
    ['AI quality', 'Correctness, relevance, groundedness, completeness, agent edit rate, task success.'],
    ['Safety and trust', 'Incorrect-action rate, policy violations, unauthorized access attempts, escalation rate, reopens.'],
    ['Operations', 'Latency, availability, provider failures, fallback rate, tool timeout rate, admin effort.'],
    ['Business', 'Pilot conversion, retention, expansion, implementation time, gross margin after model and connector cost.']
], [1.5, 4.9])
p(doc, 'Metric definition: a verified resolution is a request resolved with AI assistance, confirmed by the customer or human agent, with no avoidable reopen or escalation within a defined review window. The review window and confirmation method must be agreed with pilot customers.')

h(doc, '11 Business model hypothesis')
p(doc, 'The market contains seat, session, conversation, action, and outcome pricing. A predictable hybrid model is the safest initial hypothesis because customers need budget visibility while the company must recover model, retrieval, connector, execution, onboarding, and support costs.')
add_table(doc, ['Component', 'Hypothesis'], [
    ['Base subscription', 'Access to the control layer, administration, analytics, and standard integrations.'],
    ['Usage tier', 'Support volume, AI workload, or verified resolution band with clear limits.'],
    ['Premium package', 'Advanced governance, audit, evaluation, data residency, and enterprise controls.'],
    ['Expansion path', 'One workflow -> more workflows -> more teams -> higher governance and usage tier.'],
    ['Margin guardrail', 'Track inference, retrieval, storage, connector, execution, human review, onboarding, and support costs.']
], [1.7, 4.7])

h(doc, '12 Risks and mitigation')
add_table(doc, ['Risk', 'Why it matters', 'Mitigation'], [
    ['Incumbents copy the feature', 'Core agent capabilities are already in major platforms.', 'Differentiate through a validated cross-system workflow, not a generic agent.'],
    ['No willingness to pay for a separate layer', 'Customers may prefer an incumbent add-on.', 'Test purchase intent and compare against embedded workflow option.'],
    ['Integration work destroys margin', 'Every customer may require custom connectors.', 'Limit initial connectors and build reusable patterns.'],
    ['Poor knowledge produces poor automation', 'Grounding quality limits resolution quality.', 'Add content diagnostics, citations, freshness checks, and human fallback.'],
    ['Automation causes customer harm', 'Incorrect actions can damage trust and revenue.', 'Approval gates, allowlists, audit, rollback, and staged rollout.'],
    ['Platform scope becomes too broad', 'Customer operations can become an unbounded programme.', 'Use quarterly evidence gates and explicit non-goals.']
], [1.5, 2.2, 2.7])

h(doc, '13 Traceability matrix')
add_table(doc, ['Customer problem', 'Product response', 'Metric', 'Evidence needed'], [
    ['Agents lack customer context.', 'Identity resolution and approved context retrieval.', 'Time to resolution; correction rate.', 'Observed workflow and system map.'],
    ['Agents repeat manual steps.', 'Low-risk workflow actions and routing.', 'Verified resolutions; handle time.', 'Frequency, severity, and action permission.'],
    ['Leaders cannot trust AI outcomes.', 'Quality, audit, provenance, and governance.', 'Reopen, policy violation, escalation rate.', 'Buyer and security requirements.'],
    ['AI value is difficult to defend.', 'Outcome dashboard and baseline comparison.', 'Cost per resolved request; retention; expansion.', 'Baseline and pilot control period.']
], [1.8, 2.0, 1.5, 1.5])

doc.add_page_break()
h(doc, '14 Five-minute Loom script')
add_table(doc, ['Time', 'Narrative'], [
    ['0:00–0:25', 'The company is at risk of becoming another AI chatbot while customers continue using existing helpdesk and CRM systems.'],
    ['0:25–1:00', 'Market evidence shows incumbents already provide agents, actions, governance, and outcomes. The opportunity must be narrower.'],
    ['1:00–1:35', 'Target support-led B2B companies with growing demand, lean teams, and disconnected systems. Validate this segment through interviews.'],
    ['1:35–2:15', 'Strategic choice: become a governed AI support-operations layer that automates selected resolutions across existing systems.'],
    ['2:15–3:00', 'Start with one workflow, then add safe actions, governance, orchestration, and later broader customer operations.'],
    ['3:00–3:35', 'Use a hybrid subscription and usage model; expand from one workflow to more teams and workflows.'],
    ['3:35–4:15', 'Measure verified resolutions while protecting reopens, policy violations, escalation, latency, cost, and CSAT.'],
    ['4:15–4:45', 'Close with the unresolved decision: will customers pay for a separate vendor-neutral layer? Interviews and a controlled pilot must answer this.']
], [1.1, 5.3])

h(doc, '15 Immediate execution checklist')
for item in [
    'Recruit 6–8 interview participants across company sizes, industries, support maturity, and roles.',
    'Complete interviews using recent-workflow questions and anonymized notes.',
    'Populate the evidence log and label each insight as evidence, assumption, hypothesis, or recommendation.',
    'Update the ICP and strategic options based on interview patterns.',
    'Validate whether the cross-system problem is more urgent than improving the existing helpdesk.',
    'Define one pilot workflow and its baseline metrics.',
    'Document security, permissions, human approval, audit, fallback, and rollback requirements.',
    'Create the final written submission and record the Loom only after the evidence log is complete.'
]: bullet(doc, item)

h(doc, '16 Source evidence data')
p(doc, 'This is the extracted source data behind the strategy. These figures are not customer interviews or internal product data. Most are vendor-sponsored surveys or vendor-reported product claims, so they are directional and must be labelled accordingly in the assessment.')
add_table(doc, ['Source data point', 'Source and date', 'What it supports', 'Limitation'], [
    ['2,470 support professionals surveyed; Q4 2025; 82% of senior leaders invested in AI in the prior 12 months; 87% planned to invest in 2026; only 10% reported mature, fully integrated deployment; mature teams reporting improved metrics: 87% versus 62% overall.', 'Intercom 2026 Customer Service Transformation Report; published 2026.', 'AI adoption is widespread, but operational maturity is still limited.', 'Intercom-sponsored research; sample and methodology are defined by the vendor.'],
    ['Over 2,000 customer-service professionals surveyed; 76% invested in AI versus 54% planned; 79% planned investment in the following year; only 19% said tools always fully support their needs; 81% said AI is changing the economics of customer service.', 'Intercom 2025 Customer Service Transformation Report; published January 2025.', 'There may be a gap between AI investment and tool/workflow satisfaction.', 'Vendor-sponsored survey; does not prove a separate control-layer buying decision.'],
    ['Only 20% of organizations reported having a generative-AI governance strategy, and 23% felt highly prepared to manage it.', 'Zendesk AI Trust Gap Report; updated July 2025.', 'Governance and preparedness are plausible adoption barriers.', 'Zendesk-sponsored research; not a neutral industry benchmark.'],
    ['Research with 300 U.S. professionals evaluating customer-service platforms ranked security, compliance, and governance as the most important capability group; compliance/security was also the top operational challenge.', 'Zendesk AI Governance research; updated May 2026.', 'Trust and governance can affect the purchase decision, not only post-sale operations.', 'U.S.-only sample and vendor research.'],
    ['Security concerns were reported as the number-one AI implementation challenge; more than half said security concerns delayed or limited initiatives; 86% said they would pay more for technology that keeps data secure.', 'Salesforce State of Service, Seventh Edition.', 'Security, data protection, and unified data should be treated as product and procurement requirements.', 'Salesforce-sponsored report; willingness to pay is stated in a survey context, not validated purchasing behaviour.'],
    ['$0.99 per successful Fin outcome for listed resolution and procedure-handoff outcomes; unsuccessful attempts and explicit human escalations are not charged.', 'Intercom Fin outcome pricing; updated July 2026.', 'Outcome-based pricing is already present in the category.', 'One competitor’s pricing and outcome definition; not proof that outcome pricing is right for this product.'],
    ['Freshdesk plans list AI agents, agentic workflows, analytics, routing, permissions, and audit logs; AI sessions are priced separately on listed plans.', 'Freshworks Freshdesk features and pricing pages; current pages accessed September 2026.', 'Incumbents already combine support workflow, AI, governance, and usage packaging.', 'Product-page claims and pricing can change; feature availability may vary by plan and region.'],
    ['Salesforce Agentforce lists $0.10 per action through Flex Credits and $2 per conversation.', 'Salesforce Agentforce pricing help; published May 2025.', 'Action and conversation pricing are competitive category patterns.', 'List pricing is not a customer’s total cost of ownership.']
], [2.45, 1.55, 1.55, 0.9])

h(doc, '17 Sources reviewed')
for source in [
    'Intercom Fin AI Agent outcomes: https://www.intercom.com/help/en/articles/8205718-fin-ai-agent-outcomes',
    'Intercom Fin AI Agent explained: https://www.intercom.com/help/en/articles/7120684-fin-ai-agent-explained',
    'Zendesk AI agents: https://www.zendesk.com/service/ai/ai-agents/?lang=en',
    'Zendesk AI trust: https://www.zendesk.com/service/ai/ai-trust/',
    'Freshdesk helpdesk features: https://www.freshworks.com/freshdesk/helpdesk-features/',
    'Salesforce Agentforce pricing: https://help.salesforce.com/s/articleView?id=004811240&language=en_US&type=1',
    'Salesforce State of Service, Seventh Edition: https://www.salesforce.com/en-eu/wp-content/uploads/sites/11/documents/PDF/state-of-service-7th-edition.pdf'
]: bullet(doc, source)
p(doc, 'Source limitation: most market and competitor evidence is vendor-published. Use it as directional evidence about category positioning and product capabilities, not as independent proof of customer outcomes or market size.')

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('Airtribe B2B Product Strategy Assessment | Working evidence pack')
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

doc.save(DOCX)
print(DOCX)
